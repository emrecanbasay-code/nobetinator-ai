import io
import unittest
import openpyxl
import pandas as pd
from nobet_on_inceleme import compute,validate_reduction,export_report,fingerprint,istek_karsilastirma_satirlari


def data(docs=('A','B','C','D'),q24=4,q16=2):
    return dict(docs=list(docs),year=2026,month=2,rest=2,q24={d:q24 for d in docs},q16={d:q16 for d in docs},manual={},seniority={},couples=[])


def check(test, d, r):
    test.assertEqual(r['status'],'OK')
    test.assertEqual(len(r['rows']),28)
    for row in r['rows']:
        test.assertLessEqual(row['n16'],2)
        test.assertFalse(set(row['team24']) & set(row['team16']))
        for doc in row['team24']+row['team16']:
            test.assertNotEqual(d['manual'].get(f'{doc}_{row["day"]}'),'X')
    for doc in d['docs']:
        for i,row in enumerate(r['rows']):
            hours=24 if doc in row['team24'] else 16 if doc in row['team16'] else 0
            fixed=d['manual'].get(f'{doc}_{i+1}')
            if fixed in ('24','16'):test.assertEqual(hours,int(fixed))
            if hours:
                for later in r['rows'][i+1:i+1+(d['rest'] if hours==24 else 1)]:
                    test.assertNotIn(doc,later['team24']+later['team16'])
        test.assertLessEqual(sum(doc in row['team24'] for row in r['rows']),d['q24'][doc])
        test.assertLessEqual(sum(doc in row['team16'] for row in r['rows']),d['q16'][doc])


class Tests(unittest.TestCase):
    def test_rotation_forced_exceptions_preserve_total_and_report(self):
        d=data(('A','B'),0,0)
        d['q24']={'A':2,'B':0};d['q16']={'A':1,'B':1}
        d['rotation_doctors']=['A','B']
        d['manual']={'A_7':'24','B_7':'16','A_14':'24','A_21':'16'}
        r=compute(d,2);check(self,d,r)
        self.assertEqual(r['assigned'],4)
        self.assertEqual(len(r.get('rotation_warnings',[])),2)

    def test_rotation_weekend_preference_uses_weekday_alternative(self):
        d=data(('A',),2,1)
        d['rotation_doctors']=['A']
        d['manual']={f'A_{t}':'X' for t in range(1,29) if t not in (7,14,20,21)}
        d['manual'].update({'A_7':'24','A_14':'24'})
        r=compute(d,2);check(self,d,r)
        self.assertEqual(r['assigned'],3)
        self.assertEqual(r['rows'][19]['team16'],['A'])
        self.assertEqual(r.get('rotation_warnings'),[])

    def test_full(self):
        d=data();d['manual']={'A_1':'24','B_1':'X'}
        r=compute(d,2);check(self,d,r)
        self.assertEqual(r['assigned'],r['target'])
    def test_shortage(self):
        d=data(('A',),31,0);r=compute(d,1);check(self,d,r)
        self.assertEqual(r['assigned'],10);self.assertTrue(r['proven'])
    def test_zero(self):
        d=data(q24=0,q16=0);r=compute(d,1);check(self,d,r)
        self.assertEqual(r['assigned'],0)
    def test_quota_fixed_conflict(self):
        d=data(q24=0,q16=0);d['manual']={'A_1':'24'}
        r=compute(d,1);self.assertEqual(r['status'],'INFEASIBLE');self.assertTrue(r['reasons'])
    def test_two_fixed_sixteen(self):
        d=data();d['manual']={'A_1':'16','B_1':'16'}
        r=compute(d,1);check(self,d,r)
        self.assertEqual(r['rows'][0]['n16'],2)
    def test_three_fixed_sixteen(self):
        d=data();d['manual']={'A_1':'16','B_1':'16','C_1':'16'}
        self.assertEqual(compute(d,1)['status'],'INFEASIBLE')
    def test_fixed_rest(self):
        d=data();d['manual']={'A_26':'24','A_28':'16'}
        self.assertEqual(compute(d,1)['status'],'INFEASIBLE')
    def test_reduction(self):
        d=data();d['manual']={'A_1':'24'};r=compute(d,1)
        rows=[{'24 saat':x['n24'],'16 saat':x['n16']} for x in r['rows']]
        a,b=validate_reduction(rows,r,d['manual']);self.assertEqual(len(a),28)
        rows[0]['24 saat']=0
        with self.assertRaises(ValueError):validate_reduction(rows,r,d['manual'])
    def test_increase_rejected(self):
        d=data();r=compute(d,1)
        rows=[{'24 saat':x['n24'],'16 saat':x['n16']} for x in r['rows']]
        rows[0]['24 saat']+=1
        with self.assertRaises(ValueError):validate_reduction(rows,r,d['manual'])
    def test_istek_karsilastirma_satirlari(self):
        manual={'A_1':'X','A_2':'S','A_3':'24','A_4':'16','B_1':'24','B_2':'S'}
        duties={'A':{2:'24',3:'24',4:'16'},'B':{}}
        rows=istek_karsilastirma_satirlari(['A','B'],manual,duties,2026,10)
        self.assertEqual(rows[0]['Personel'],'A');self.assertEqual(rows[1]['Personel'],'B')
        self.assertEqual(rows[0]['01.10'],'X')
        self.assertEqual(rows[0]['02.10'],'S→24')
        self.assertEqual(rows[0]['03.10'],'24')
        self.assertEqual(rows[0]['04.10'],'16')
        self.assertEqual(rows[0]['31.10'],'')
        self.assertEqual(rows[1]['01.10'],'24!')
        self.assertEqual(rows[1]['02.10'],'S')
    def test_excel_export(self):
        d=data();d['manual']={'A_1':'24','B_1':'X','C_1':'S'}
        r=compute(d,1)
        months=['','Ocak','Şubat','Mart','Nisan','Mayıs','Haziran','Temmuz','Ağustos','Eylül','Ekim','Kasım','Aralık']
        excel=pd.ExcelFile(io.BytesIO(export_report(d,r,months)))
        self.assertEqual(excel.sheet_names,['Ozet','Gunluk Dagilim','Kisi Bazinda','Kisi Nobetleri','Istek Karsilastirma','Ornek Cizelge'])
        self.assertEqual(len(pd.read_excel(excel,sheet_name='Gunluk Dagilim')),28)
        self.assertEqual(len(pd.read_excel(excel,sheet_name='Kisi Bazinda')),4)
        kisiler=pd.read_excel(excel,sheet_name='Kisi Nobetleri')
        gunler=[c for c in kisiler.columns if c not in ('Personel','Toplam')]
        self.assertEqual(gunler[0],'01.02');self.assertEqual(gunler[-1],'28.02')
        # Hücre içerikleri pandas okumasında sayıya dönüşebildiğinden dosyadan doğrulanır.
        wb=openpyxl.load_workbook(io.BytesIO(export_report(d,r,months)))
        ws=wb['Kisi Nobetleri']
        self.assertEqual(ws.max_row,5);self.assertEqual(ws.max_column,30)
        self.assertEqual(ws.cell(1,1).value,'Personel')
        self.assertEqual(ws.cell(1,2).value,'01.02')
        self.assertEqual(ws.cell(1,29).value,'28.02')
        self.assertEqual(ws.cell(1,30).value,'Toplam')
        for satir in range(2,ws.max_row+1):
            doktor=ws.cell(satir,1).value
            istat=next(s for s in r['stats'] if s['Doktor']==doktor)
            degerler=[ws.cell(satir,c).value for c in range(2,30)]
            self.assertTrue(all(v in ('24','16',None) for v in degerler))
            self.assertEqual(degerler.count('24'),istat['24s yazılan'])
            self.assertEqual(degerler.count('16'),istat['16s yazılan'])
            self.assertEqual(ws.cell(satir,30).value,istat['24s yazılan']+istat['16s yazılan'])
        kars=wb['Istek Karsilastirma']
        self.assertEqual(kars.cell(1,1).value,'Personel')
        self.assertEqual(kars.cell(1,2).value,'01.02')
        self.assertEqual(kars.cell(2,2).value,'24')
        self.assertEqual(kars.cell(3,2).value,'X')
        self.assertIn(kars.cell(4,2).value,('S','S→24','S→16'))
    def test_excel_export_includes_warnings(self):
        d=data()
        r={'status':'OK','rows':[{'day':1,'date':'2026-02-01','n24':1,'n16':2,'team24':['A'],'team16':['B','C']}],
           'stats':[{'Doktor':'A','24s hedef':1,'24s yazılan':1,'24s eksik':0,'16s hedef':0,'16s yazılan':0,'16s eksik':0}],
           'target':3,'assigned':3,'proven':True,'violations':['A: 1. gün 24 saat'],
           'rotation_warnings':['B: örnek uyarı'],'incoming_warnings':['C: gelen uyarısı']}
        excel=pd.ExcelFile(io.BytesIO(export_report(d,r,['']*13)))
        self.assertIn('Uyarilar',excel.sheet_names)
        self.assertEqual(pd.read_excel(excel,sheet_name='Uyarilar')['Uyarılar'].tolist(),
                         ['S izni: A: 1. gün 24 saat','B: örnek uyarı','C: gelen uyarısı'])
        self.assertEqual(pd.read_excel(excel,sheet_name='Gunluk Dagilim')['Toplam'].tolist(),[3])
    def test_changed_input(self):
        d=data();a=fingerprint(d);d['q24']['A']=7;self.assertNotEqual(a,fingerprint(d))
    def test_incoming_spread_across_days(self):
        d=data();d['incoming']=['A','B']
        r=compute(d,3);check(self,d,r)
        days_a={row['day'] for row in r['rows'] if 'A' in row['team24']+row['team16']}
        days_b={row['day'] for row in r['rows'] if 'B' in row['team24']+row['team16']}
        self.assertFalse(days_a & days_b)
        self.assertEqual(r.get('incoming_warnings'),[])
    def test_incoming_forced_same_day_reported(self):
        d=data(('A','B'),0,0)
        d['q24']={'A':1,'B':1};d['q16']={'A':0,'B':0};d['incoming']=['A','B']
        d['manual']={'A_7':'24','B_7':'24'}
        r=compute(d,2);check(self,d,r)
        self.assertEqual(r['assigned'],2)
        self.assertEqual(len(r['incoming_warnings']),1)
        self.assertIn('Rotasyona gelenler',r['incoming_warnings'][0])
        self.assertIn('07.02.2026',r['incoming_warnings'][0])
    def test_all_blocked(self):
        d=data(('A',),4,2);d['manual']={f'A_{t}':'X' for t in range(1,29)}
        r=compute(d,1);check(self,d,r);self.assertEqual(r['assigned'],0)

if __name__=='__main__':unittest.main()
